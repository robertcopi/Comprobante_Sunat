import csv
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse
from django.shortcuts import redirect
from django.urls import reverse_lazy
from django.views import View
from django.views.generic import ListView, DetailView, CreateView
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from apps.accounts.views import get_client_ip
from .models import LoteImportacion
from .forms import ImportacionForm
from .services import ImportacionService


class ImportacionListView(LoginRequiredMixin, ListView):
    """
    Listado histórico de lotes de importación masiva.
    """
    model = LoteImportacion
    template_name = 'importaciones/importacion_list.html'
    context_object_name = 'lotes'
    paginate_by = 10


class ImportacionUploadView(LoginRequiredMixin, CreateView):
    """
    Subida y procesamiento masivo de archivos Excel o CSV.
    Permiso: Rol TRABAJADOR o ADMINISTRADOR.
    """
    model = LoteImportacion
    form_class = ImportacionForm
    template_name = 'importaciones/importacion_form.html'

    def dispatch(self, request, *args, **kwargs):
        user = request.user
        if not (user.is_superuser or user.role in ['ADMINISTRADOR', 'TRABAJADOR']):
            messages.error(request, "Acceso restringido: El rol Supervisor no cuenta con permisos para importar archivos.")
            return redirect('importacion_list')
        return super().dispatch(request, *args, **kwargs)

    def form_valid(self, form):
        archivo = form.cleaned_data['archivo']
        lote = form.save(commit=False)
        lote.nombre_archivo = archivo.name
        lote.usuario = self.request.user
        lote.save()

        # Ejecutar servicio de procesamiento y validación
        service = ImportacionService(lote, ip_address=get_client_ip(self.request))
        lote_procesado = service.procesar()

        messages.success(
            self.request,
            f"Archivo '{lote.nombre_archivo}' procesado: {lote.importados_correctos} comprobantes importados exitosamente "
            f"({lote.total_duplicados} duplicados, {lote.total_errores} con error)."
        )
        return redirect('importacion_detail', pk=lote.pk)


class ImportacionDetailView(LoginRequiredMixin, DetailView):
    """
    Vista detallada del resultado de un lote de importación, métricas y detalle de incidencias.
    """
    model = LoteImportacion
    template_name = 'importaciones/importacion_detail.html'
    context_object_name = 'lote'


class DescargarPlantillaExcelView(LoginRequiredMixin, View):
    """
    Genera y descarga un archivo de plantilla Excel (.xlsx) con el formato oficial esperado.
    """
    def get(self, request, *args, **kwargs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comprobantes"

        headers = ['RUC', 'TIPO', 'SERIE', 'NUMERO', 'FECHA', 'MONTO']
        ws.append(headers)

        # Estilo encabezado
        header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Filas de ejemplo
        ejemplos = [
            ['20100070970', '03', 'B001', 101, '2026-09-20', 85.50],
            ['20100070970', '03', 'B001', 102, '2026-09-20', 140.00],
            ['20555666777', '01', 'F001', 500, '2026-09-21', 1250.00],
        ]
        for ej in ejemplos:
            ws.append(ej)

        # Ajustar ancho columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="plantilla_comprobantes_sunat.xlsx"'
        wb.save(response)
        return response


class DescargarPlantillaCsvView(LoginRequiredMixin, View):
    """
    Genera y descarga un archivo de plantilla CSV (.csv) con el formato oficial esperado.
    """
    def get(self, request, *args, **kwargs):
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="plantilla_comprobantes_sunat.csv"'

        writer = csv.writer(response)
        writer.writerow(['RUC', 'TIPO', 'SERIE', 'NUMERO', 'FECHA', 'MONTO'])
        writer.writerow(['20100070970', '03', 'B001', '101', '2026-09-20', '85.50'])
        writer.writerow(['20100070970', '03', 'B001', '102', '2026-09-20', '140.00'])
        writer.writerow(['20555666777', '01', 'F001', '500', '2026-09-21', '1250.00'])

        return response
