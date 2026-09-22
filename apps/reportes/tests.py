import io
from datetime import date, timedelta
from openpyxl import load_workbook
from django.contrib.auth import get_user_model
from django.test import TestCase, Client
from django.urls import reverse

from apps.auditoria.models import RegistroAuditoria
from apps.comprobantes.models import Comprobante
from apps.validaciones.models import ValidacionSunat

User = get_user_model()


class ReportesModuloTests(TestCase):
    """
    Pruebas integrales para el módulo de Reportes (PASO 14):
    - Filtrado por los 7 criterios (fecha inicio, fecha fin, RUC, estado, tipo, serie, usuario).
    - Visualización en pantalla con datos agregados.
    - Exportación a Excel (.xlsx) con las 10 columnas obligatorias.
    - Exportación a PDF corporativo.
    - Cumplimiento de filtros en la exportación.
    - Registro de cada exportación en la bitácora de auditoría.
    """

    def setUp(self):
        self.client = Client()
        self.user1 = User.objects.create_user(
            username='supervisor_reportes',
            password='SupervisorPass123!',
            role='SUPERVISOR'
        )
        self.user2 = User.objects.create_user(
            username='operador_dos',
            password='OperadorPass123!',
            role='TRABAJADOR'
        )

        hoy = date.today()
        hace_3_dias = hoy - timedelta(days=3)
        hace_10_dias = hoy - timedelta(days=10)

        # Comprobante 1: Factura Aceptada
        self.cp1 = Comprobante.objects.create(
            ruc_emisor='20111111111',
            tipo_comprobante='01',
            serie='F001',
            numero=1001,
            fecha_emision=hoy,
            monto=500.00,
            estado=Comprobante.EstadoComprobante.ACEPTADO,
            usuario_registro=self.user1
        )
        self.val1 = ValidacionSunat.objects.create(
            comprobante=self.cp1,
            estado_sunat='Aceptado',
            codigo_respuesta='1',
            mensaje='Comprobante Aceptado por SUNAT',
            usuario=self.user1
        )

        # Comprobante 2: Boleta Observada
        self.cp2 = Comprobante.objects.create(
            ruc_emisor='20222222222',
            tipo_comprobante='03',
            serie='B001',
            numero=2002,
            fecha_emision=hace_3_dias,
            monto=150.00,
            estado=Comprobante.EstadoComprobante.OBSERVADO,
            usuario_registro=self.user2
        )
        self.val2 = ValidacionSunat.objects.create(
            comprobante=self.cp2,
            estado_sunat='Observado',
            codigo_respuesta='1',
            mensaje='Emisor con condición de No Habido',
            usuario=self.user2
        )

        # Comprobante 3: Factura Rechazada
        self.cp3 = Comprobante.objects.create(
            ruc_emisor='20333333333',
            tipo_comprobante='01',
            serie='F002',
            numero=3003,
            fecha_emision=hace_10_dias,
            monto=1200.00,
            estado=Comprobante.EstadoComprobante.RECHAZADO,
            usuario_registro=self.user1
        )

    def test_reporte_requiere_login(self):
        resp = self.client.get(reverse('reporte_comprobantes'))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse('login'), resp.url)

    def test_reporte_list_view_renderiza_pantalla(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        resp = self.client.get(reverse('reporte_comprobantes'))
        self.assertEqual(resp.status_code, 200)

        ctx = resp.context
        self.assertEqual(ctx['total_registros'], 3)
        self.assertEqual(ctx['monto_total'], 1850.00)
        self.assertEqual(ctx['total_aceptados'], 1)
        self.assertEqual(ctx['total_observados'], 1)
        self.assertEqual(ctx['total_rechazados'], 1)

        self.assertContains(resp, "F001-00001001")
        self.assertContains(resp, "B001-00002002")
        self.assertContains(resp, "F002-00003003")

    def test_filtro_por_fecha_inicio_y_fin(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        hoy = date.today()
        hace_5_dias = hoy - timedelta(days=5)

        # Filtro: últimos 5 días (debe excluir cp3 de hace 10 días)
        url = reverse('reporte_comprobantes') + f'?fecha_inicio={hace_5_dias}&fecha_fin={hoy}'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)

        ctx = resp.context
        self.assertEqual(ctx['total_registros'], 2)
        codigos = [cp.codigo_completo for cp in ctx['comprobantes']]
        self.assertIn("F001-00001001", codigos)
        self.assertIn("B001-00002002", codigos)
        self.assertNotIn("F002-00003003", codigos)

    def test_filtro_por_ruc(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_comprobantes') + '?ruc=20222222222'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['total_registros'], 1)
        self.assertEqual(resp.context['comprobantes'][0].codigo_completo, "B001-00002002")

    def test_filtro_por_estado(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_comprobantes') + '?estado=ACEPTADO'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['total_registros'], 1)
        self.assertEqual(resp.context['comprobantes'][0].estado, 'ACEPTADO')

    def test_filtro_por_tipo_comprobante(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_comprobantes') + '?tipo_comprobante=01'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        # cp1 y cp3 son facturas (01)
        self.assertEqual(resp.context['total_registros'], 2)

    def test_filtro_por_serie(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_comprobantes') + '?serie=B001'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['total_registros'], 1)
        self.assertEqual(resp.context['comprobantes'][0].serie, 'B001')

    def test_filtro_por_usuario(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_comprobantes') + f'?usuario={self.user2.id}'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context['total_registros'], 1)
        self.assertEqual(resp.context['comprobantes'][0].codigo_completo, "B001-00002002")

    def test_exportar_excel_contiene_10_columnas_requeridas(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_exportar_excel')
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('attachment; filename="reporte_comprobantes_', resp['Content-Disposition'])

        # Cargar libro con openpyxl y verificar las 10 columnas obligatorias
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active

        # Fila 2 contiene los encabezados
        columnas_esperadas = [
            'RUC',
            'TIPO',
            'SERIE',
            'NUMERO',
            'FECHA EMISION',
            'MONTO',
            'ESTADO',
            'FECHA VALIDACION',
            'MENSAJE',
            'USUARIO'
        ]
        columnas_obtenidas = [ws.cell(row=2, column=i).value for i in range(1, 11)]
        self.assertEqual(columnas_obtenidas, columnas_esperadas)

        # Debe contener 3 comprobantes (filas 3, 4, 5)
        self.assertIsNotNone(ws.cell(row=3, column=1).value)
        self.assertIsNotNone(ws.cell(row=4, column=1).value)
        self.assertIsNotNone(ws.cell(row=5, column=1).value)

        # Verificar que se registró en la bitácora de auditoría
        audit = RegistroAuditoria.objects.filter(accion='EXPORTAR_REPORTE', id_objeto='reporte_excel').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.usuario, self.user1)
        self.assertEqual(audit.resultado, 'EXITOSO')

    def test_exportar_pdf_genera_documento_y_auditoria(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        url = reverse('reporte_exportar_pdf')
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp['Content-Type'], 'application/pdf')
        self.assertIn('attachment; filename="reporte_comprobantes_', resp['Content-Disposition'])
        self.assertTrue(resp.content.startswith(b'%PDF'))

        # Verificar auditoría
        audit = RegistroAuditoria.objects.filter(accion='EXPORTAR_REPORTE', id_objeto='reporte_pdf').first()
        self.assertIsNotNone(audit)
        self.assertEqual(audit.usuario, self.user1)
        self.assertEqual(audit.resultado, 'EXITOSO')

    def test_exportacion_respeta_filtros_activos(self):
        self.client.login(username='supervisor_reportes', password='SupervisorPass123!')
        # Exportar filtrando solo por estado ACEPTADO
        url = reverse('reporte_exportar_excel') + '?estado=ACEPTADO'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)

        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active

        # Solo debe haber 1 comprobante en fila 3
        self.assertEqual(ws.cell(row=3, column=1).value, '20111111111')
        self.assertEqual(ws.cell(row=3, column=7).value, 'ACEPTADO')
        # Fila 4 debe estar vacía
        self.assertIsNone(ws.cell(row=4, column=1).value)
